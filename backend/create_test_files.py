from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import json
from datetime import datetime

# ============================================================
# 1. EXCEL MODÈLE (foana)
# ============================================================
wb = Workbook()
ws = wb.active
ws.title = "Eleves"

headers = ["matricule", "nom", "prenom", "date_naissance", "age", "sexe",
           "lieu", "parent_nom", "parent_telephone", "telephone", "classe"]
ws.append(headers)

# Style des en-têtes
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
for col_num, _ in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Largeur des colonnes
widths = [15, 15, 15, 15, 6, 6, 18, 18, 18, 15, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + i)].width = w

# Ligne d'exemple
ws.append(["", "Rakoto", "Jean", "2010-05-15", 15, "M",
           "Antananarivo", "Rabe", "0341234567", "0341111111", "Classe A"])

wb.save("modele_eleves.xlsx")
print("✅ modele_eleves.xlsx noforonina")

# ============================================================
# 2. EXCEL MISY DONNÉES DE TEST
# ============================================================
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Eleves"
ws2.append(headers)
for col_num, _ in enumerate(headers, 1):
    cell = ws2.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")
for i, w in enumerate(widths, 1):
    ws2.column_dimensions[chr(64 + i)].width = w

# 30 élèves de test
eleves = [
    ["", "RAKOTO", "Jean", "2010-05-15", 15, "M", "Antananarivo", "RABE Jean", "0341234567", "0341111111", "Classe A"],
    ["", "RANDRIA", "Marie", "2011-03-22", 14, "F", "Antananarivo", "RANDRIA Paul", "0342345678", "0342222222", "Classe A"],
    ["", "RAZAFY", "Pierre", "2010-08-10", 15, "M", "Antsirabe", "RAZAFY Luc", "0343456789", "0343333333", "Classe A"],
    ["", "ANDRIANINA", "Sophie", "2011-01-18", 14, "F", "Fianarantsoa", "ANDRIANINA Marc", "0344567890", "0344444444", "Classe A"],
    ["", "RAKOTOMALALA", "Paul", "2010-11-30", 14, "M", "Toliara", "RAKOTOMALALA Henri", "0345678901", "0345555555", "Classe A"],
    ["", "RASOAMANANA", "Hery", "2011-07-04", 14, "M", "Antananarivo", "RASOAMANANA Guy", "0346789012", "0346666666", "Classe B"],
    ["", "RAHARISON", "Lalao", "2010-09-12", 15, "F", "Mahajanga", "RAHARISON Albert", "0347890123", "0347777777", "Classe B"],
    ["", "RAKOTOARISOA", "Fidy", "2011-02-28", 14, "M", "Antananarivo", "RAKOTOARISOA Bruno", "0348901234", "0348888888", "Classe B"],
    ["", "RANDRIANASOLO", "Tiana", "2010-06-20", 15, "F", "Antananarivo", "RANDRIANASOLO Eric", "0349012345", "0349999999", "Classe B"],
    ["", "RAVELO", "Njaka", "2011-04-15", 14, "M", "Antsirabe", "RAVELO Jean", "0340123456", "0340000000", "Classe B"],
    ["", "RAZANADRAKOTO", "Vola", "2010-10-25", 14, "F", "Toamasina", "RAZANADRAKOTO Hery", "0341111112", "0341122334", "Classe C"],
    ["", "RABEMANANJARA", "Tahiry", "2011-05-08", 14, "M", "Antananarivo", "RABEMANANJARA Solofo", "0341222223", "0342233445", "Classe C"],
    ["", "RANDRIAMAMPIONONA", "Fara", "2010-12-03", 14, "F", "Antananarivo", "RANDRIAMAMPIONONA Guy", "0341333334", "0343344556", "Classe C"],
    ["", "RAKOTONDRABE", "Mamy", "2011-08-19", 14, "M", "Fianarantsoa", "RAKOTONDRABE Luc", "0341444445", "0344455667", "Classe C"],
    ["", "RASOLOFOSON", "Andry", "2010-07-11", 15, "M", "Antananarivo", "RASOLOFOSON Paul", "0341555556", "0345566778", "Classe C"],
    ["", "RAJAONARIVELO", "Fenitra", "2011-11-27", 13, "F", "Antananarivo", "RAJAONARIVELO Bruno", "0341666667", "0346677889", "Classe A"],
    ["", "RAKOTOVAO", "Tovo", "2010-02-14", 15, "M", "Toliara", "RAKOTOVAO Henri", "0341777778", "0347788990", "Classe A"],
    ["", "RANOROSOA", "Lova", "2011-09-30", 14, "F", "Antsirabe", "RANOROSOA Marc", "0341888889", "0348899001", "Classe A"],
    ["", "RABE", "Tsiory", "2010-04-07", 15, "M", "Antananarivo", "RABE Eric", "0341999990", "0349900112", "Classe B"],
    ["", "RATSIMBAZAFY", "Nomena", "2011-06-16", 14, "F", "Mahajanga", "RATSIMBAZAFY Albert", "0342000001", "0340011223", "Classe B"],
    ["", "RANDRIANARISOA", "Fandresena", "2010-03-09", 15, "M", "Antananarivo", "RANDRIANARISOA Guy", "0342111112", "0341122335", "Classe B"],
    ["", "RAHERIMANANA", "Tantely", "2011-10-21", 14, "F", "Antananarivo", "RAHERIMANANA Solofo", "0342222223", "0342233446", "Classe C"],
    ["", "RAKOTOMANGA", "Fy", "2010-01-05", 15, "M", "Fianarantsoa", "RAKOTOMANGA Bruno", "0342333334", "0343344557", "Classe C"],
    ["", "RASOANAIVO", "Miora", "2011-12-12", 13, "F", "Antananarivo", "RASOANAIVO Luc", "0342444445", "0344455668", "Classe C"],
    ["", "RANDRIANIRINA", "Feno", "2010-08-28", 15, "M", "Toamasina", "RANDRIANIRINA Paul", "0342555556", "0345566779", "Classe A"],
    ["", "RAKOTOMAVO", "Hanta", "2011-02-02", 14, "F", "Antananarivo", "RAKOTOMAVO Henri", "0342666667", "0346677890", "Classe A"],
    ["", "RAZAFIMAHATRATRA", "Fidy", "2010-05-25", 15, "M", "Antsirabe", "RAZAFIMAHATRATRA Marc", "0342777778", "0347788991", "Classe B"],
    ["", "RASOLOFO", "Fanja", "2011-07-17", 14, "F", "Antananarivo", "RASOLOFO Eric", "0342888889", "0348899002", "Classe B"],
    ["", "RAKOTOBE", "Tiana", "2010-11-08", 14, "M", "Mahajanga", "RAKOTOBE Albert", "0342999990", "0349900113", "Classe C"],
    ["", "RANDRIA", "Soa", "2011-03-30", 14, "F", "Antananarivo", "RANDRIA Guy", "0343000001", "0340011224", "Classe C"],
]

for e in eleves:
    ws2.append(e)

wb2.save("eleves_test.xlsx")
print("✅ eleves_test.xlsx noforonina (30 élèves)")

# ============================================================
# 3. FICHIER JSON — BACKUP misy données
# ============================================================
backup = {
    "version": "1.0",
    "date": datetime.now().isoformat(),
    "classes": [
        {"id": 1, "nom": "Classe A", "description": "", "created_at": datetime.now().isoformat()},
        {"id": 2, "nom": "Classe B", "description": "", "created_at": datetime.now().isoformat()},
        {"id": 3, "nom": "Classe C", "description": "", "created_at": datetime.now().isoformat()},
    ],
    "groupes": [
        {"id": 1, "nom": "Groupe 1", "classe_id": 1, "created_at": datetime.now().isoformat()},
        {"id": 2, "nom": "Groupe 2", "classe_id": 1, "created_at": datetime.now().isoformat()},
        {"id": 3, "nom": "Groupe 1", "classe_id": 2, "created_at": datetime.now().isoformat()},
        {"id": 4, "nom": "Groupe 1", "classe_id": 3, "created_at": datetime.now().isoformat()},
    ],
    "students": [],
    "attendance": [],
    "sessions": [],
}

# Étudiants JSON (mifanaraka amin'ny excel)
for i, e in enumerate(eleves, 1):
    matricule = f"ELV-{datetime.now().year}-{i:04d}"
    classe_nom = e[10]
    classe_id = {"Classe A": 1, "Classe B": 2, "Classe C": 3}.get(classe_nom, 1)
    backup["students"].append({
        "id": i,
        "matricule": matricule,
        "nom": e[1],
        "prenom": e[2],
        "date_naissance": e[3],
        "age": e[4],
        "sexe": e[5],
        "lieu": e[6],
        "telephone": e[9],
        "parent_nom": e[7],
        "parent_telephone": e[8],
        "classe_id": classe_id,
        "groupe_id": None,
        "situation": "",
        "photo": "",
        "date_inscription": datetime.now().isoformat(),
        "created_at": datetime.now().isoformat(),
        "updated_at": datetime.now().isoformat(),
    })

# Sessions de test (7 andro lasa)
from datetime import timedelta
for d in range(7, 0, -1):
    day = (datetime.now() - timedelta(days=d)).strftime("%Y-%m-%d")
    backup["sessions"].append({
        "id": d,
        "classe_id": 1,
        "groupe_id": None,
        "date": day,
        "heure_debut": "08:00",
        "heure_fin": "12:00",
        "statut": "closed",
        "created_at": day,
    })

# Présences aléatoires (5 andro farany)
import random
random.seed(42)
att_id = 1
for d in range(5, 0, -1):
    day = (datetime.now() - timedelta(days=d)).strftime("%Y-%m-%d")
    for s in backup["students"][:10]:  # 10 voalohany ihany
        statut = random.choice(["present", "present", "present", "present", "absent", "retard"])
        hour = random.randint(7, 9)
        minute = random.randint(0, 59)
        backup["attendance"].append({
            "id": att_id,
            "student_id": s["id"],
            "session_id": d,
            "date": day,
            "heure_arrivee": f"{hour:02d}:{minute:02d}:00",
            "heure_depart": None,
            "statut": statut,
            "methode": "qr",
            "justif_motif": "",
            "justif_date": None,
            "created_at": day,
        })
        att_id += 1

with open("backup_test.json", "w", encoding="utf-8") as f:
    json.dump(backup, f, ensure_ascii=False, indent=2)

print("✅ backup_test.json noforonina")
print(f"   - {len(backup['students'])} élèves")
print(f"   - {len(backup['attendance'])} présences")
print(f"   - {len(backup['sessions'])} séances")

print("\n📁 Fichiers vita:")
print("   - modele_eleves.xlsx")
print("   - eleves_test.xlsx")
print("   - backup_test.json")
