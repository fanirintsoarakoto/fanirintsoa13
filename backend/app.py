from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
import sqlite3, hashlib, secrets, io, os
from datetime import datetime, date, timedelta
from openpyxl import Workbook, load_workbook
import qrcode

app = Flask(__name__)
CORS(app)
DB = "fanirintsoa.db"

# ============================================================
# DATABASE
# ============================================================
def connect():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c

def init_db():
    c = connect()
    c.execute("""CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        email TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        role TEXT DEFAULT 'admin',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS sessions_auth (
        token TEXT PRIMARY KEY, user_id INTEGER NOT NULL,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS classes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nom TEXT UNIQUE NOT NULL,
        description TEXT DEFAULT '',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS groupes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nom TEXT NOT NULL, classe_id INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        matricule TEXT UNIQUE NOT NULL,
        nom TEXT NOT NULL, prenom TEXT NOT NULL,
        date_naissance TEXT, age INTEGER,
        sexe TEXT DEFAULT 'M', lieu TEXT DEFAULT '',
        telephone TEXT DEFAULT '',
        parent_nom TEXT DEFAULT '', parent_telephone TEXT DEFAULT '',
        classe_id INTEGER, groupe_id INTEGER,
        situation TEXT DEFAULT '', photo TEXT DEFAULT '',
        date_inscription TEXT DEFAULT CURRENT_TIMESTAMP,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        updated_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS attendance_sessions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        classe_id INTEGER, groupe_id INTEGER,
        date TEXT NOT NULL, heure_debut TEXT, heure_fin TEXT,
        statut TEXT DEFAULT 'open',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS attendance (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL, session_id INTEGER,
        date TEXT NOT NULL, heure_arrivee TEXT, heure_depart TEXT,
        statut TEXT DEFAULT 'present', methode TEXT DEFAULT 'manual',
        justif_motif TEXT DEFAULT '', justif_date TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS matieres (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nom TEXT NOT NULL,
        coefficient INTEGER DEFAULT 1,
        classe_id INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS notes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        matiere_id INTEGER NOT NULL,
        note REAL NOT NULL,
        type TEXT DEFAULT 'devoir',
        trimestre INTEGER DEFAULT 1,
        date TEXT DEFAULT CURRENT_DATE,
        remarque TEXT DEFAULT '',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
    for n in ["Classe A", "Classe B", "Classe C"]:
        try:
            c.execute("INSERT INTO classes (nom) VALUES (?)", (n,))
        except sqlite3.IntegrityError:
            pass
    c.commit()
    c.close()

init_db()

# ============================================================
# HELPERS
# ============================================================
def hash_password(pw):
    salt = secrets.token_hex(16)
    h = hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 100_000)
    return f"{salt}${h.hex()}"

def verify_password(pw, stored):
    try:
        salt, h = stored.split("$", 1)
    except ValueError:
        return False
    check = hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 100_000)
    return secrets.compare_digest(check.hex(), h)

def get_user(req):
    auth = req.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    token = auth[7:]
    c = connect()
    row = c.execute("SELECT user_id FROM sessions_auth WHERE token=?", (token,)).fetchone()
    c.close()
    return row["user_id"] if row else None

def require_auth(req):
    uid = get_user(req)
    if not uid:
        return None, (jsonify({"error": "Tsy niditra"}), 401)
    return uid, None

def gen_matricule():
    c = connect()
    year = datetime.now().year
    row = c.execute("SELECT matricule FROM students WHERE matricule LIKE ? ORDER BY id DESC LIMIT 1",
                    (f"ELV-{year}-%",)).fetchone()
    if row:
        try:
            n = int(row["matricule"].split("-")[-1]) + 1
        except:
            n = 1
    else:
        n = 1
    c.close()
    return f"ELV-{year}-{n:04d}"

def compute_age(dn):
    if not dn: return None
    try:
        d = datetime.strptime(dn, "%Y-%m-%d")
        t = date.today()
        return t.year - d.year - ((t.month, t.day) < (d.month, d.day))
    except:
        return None

# ============================================================
# ROOT
# ============================================================
@app.route("/")
def home():
    return jsonify({"message": "Fanirintsoa Scolaire API", "status": "ok", "version": "2.0"})

# ============================================================
# AUTH
# ============================================================
@app.route("/auth/register", methods=["POST"])
def register():
    d = request.json
    u = (d.get("username") or "").strip()
    e = (d.get("email") or "").strip()
    p = d.get("password") or ""
    if not u or not e or not p:
        return jsonify({"error": "Mila anarana, email, teny miafina"}), 400
    if len(p) < 4:
        return jsonify({"error": "Fohy loatra ny teny miafina"}), 400
    c = connect()
    if c.execute("SELECT id FROM users WHERE username=?", (u,)).fetchone():
        c.close(); return jsonify({"error": "Efa misy io anarana io"}), 400
    if c.execute("SELECT id FROM users WHERE email=?", (e,)).fetchone():
        c.close(); return jsonify({"error": "Efa misy io email io"}), 400
    cur = c.execute("INSERT INTO users (username,email,password_hash,role) VALUES (?,?,?,'admin')",
                    (u, e, hash_password(p)))
    uid = cur.lastrowid
    token = secrets.token_urlsafe(32)
    c.execute("INSERT INTO sessions_auth (token,user_id) VALUES (?,?)", (token, uid))
    c.commit()
    c.close()
    return jsonify({"token": token,
                    "user": {"id": uid, "username": u, "email": e, "role": "admin"}}), 201

@app.route("/auth/login", methods=["POST"])
def login():
    d = request.json
    u = (d.get("username") or "").strip()
    p = d.get("password") or ""
    c = connect()
    user = c.execute("SELECT * FROM users WHERE username=?", (u,)).fetchone()
    if not user or not verify_password(p, user["password_hash"]):
        c.close(); return jsonify({"error": "Anarana na teny miafina diso"}), 401
    token = secrets.token_urlsafe(32)
    c.execute("INSERT INTO sessions_auth (token,user_id) VALUES (?,?)", (token, user["id"]))
    c.commit()
    c.close()
    return jsonify({"token": token,
                    "user": {"id": user["id"], "username": user["username"],
                             "email": user["email"], "role": user["role"]}})

@app.route("/auth/me")
def me():
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    u = c.execute("SELECT id,username,email,role FROM users WHERE id=?", (uid,)).fetchone()
    c.close()
    return jsonify(dict(u)) if u else (jsonify({"error": "Tsy hita"}), 404)

@app.route("/auth/logout", methods=["POST"])
def logout():
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        c = connect()
        c.execute("DELETE FROM sessions_auth WHERE token=?", (auth[7:],))
        c.commit()
        c.close()
    return jsonify({"ok": True})

# ============================================================
# CLASSES
# ============================================================
@app.route("/classes")
def list_classes():
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    rows = c.execute("SELECT * FROM classes ORDER BY nom").fetchall()
    result = []
    for r in rows:
        d = dict(r)
        d["nb_students"] = c.execute("SELECT COUNT(*) FROM students WHERE classe_id=?", (r["id"],)).fetchone()[0]
        d["nb_groupes"] = c.execute("SELECT COUNT(*) FROM groupes WHERE classe_id=?", (r["id"],)).fetchone()[0]
        result.append(d)
    c.close()
    return jsonify(result)

@app.route("/classes", methods=["POST"])
def add_classe():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    n = (d.get("nom") or "").strip()
    if not n: return jsonify({"error": "Mila anarana"}), 400
    c = connect()
    try:
        cur = c.execute("INSERT INTO classes (nom, description) VALUES (?,?)",
                        (n, d.get("description", "")))
        c.commit()
        row = c.execute("SELECT * FROM classes WHERE id=?", (cur.lastrowid,)).fetchone()
        c.close()
        return jsonify(dict(row)), 201
    except sqlite3.IntegrityError:
        c.close(); return jsonify({"error": "Efa misy"}), 400

@app.route("/classes/<int:cid>", methods=["DELETE"])
def del_classe(cid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    c.execute("DELETE FROM classes WHERE id=?", (cid,))
    c.commit()
    c.close()
    return jsonify({"ok": True})

# ============================================================
# GROUPES
# ============================================================
@app.route("/groupes")
def list_groupes():
    uid, err = require_auth(request)
    if err: return err
    cid = request.args.get("classe_id")
    c = connect()
    if cid:
        rows = c.execute("SELECT * FROM groupes WHERE classe_id=? ORDER BY nom", (cid,)).fetchall()
    else:
        rows = c.execute("SELECT * FROM groupes ORDER BY nom").fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/groupes", methods=["POST"])
def add_groupe():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    n = (d.get("nom") or "").strip()
    if not n: return jsonify({"error": "Mila anarana"}), 400
    c = connect()
    cur = c.execute("INSERT INTO groupes (nom, classe_id) VALUES (?,?)",
                    (n, d.get("classe_id")))
    c.commit()
    row = c.execute("SELECT * FROM groupes WHERE id=?", (cur.lastrowid,)).fetchone()
    c.close()
    return jsonify(dict(row)), 201

@app.route("/groupes/<int:gid>", methods=["DELETE"])
def del_groupe(gid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    c.execute("DELETE FROM groupes WHERE id=?", (gid,))
    c.commit()
    c.close()
    return jsonify({"ok": True})

# ============================================================
# STUDENTS
# ============================================================
def student_full(c, sid):
    row = c.execute("""SELECT s.*, cl.nom AS classe_nom, g.nom AS groupe_nom
                       FROM students s
                       LEFT JOIN classes cl ON cl.id = s.classe_id
                       LEFT JOIN groupes g ON g.id = s.groupe_id
                       WHERE s.id=?""", (sid,)).fetchone()
    return dict(row) if row else None

@app.route("/students")
def list_students():
    uid, err = require_auth(request)
    if err: return err
    a = request.args
    q = """SELECT s.*, cl.nom AS classe_nom, g.nom AS groupe_nom
           FROM students s
           LEFT JOIN classes cl ON cl.id = s.classe_id
           LEFT JOIN groupes g ON g.id = s.groupe_id WHERE 1=1"""
    p = []
    if a.get("search"):
        s = f"%{a.get('search')}%"
        q += " AND (s.nom LIKE ? OR s.prenom LIKE ? OR s.matricule LIKE ?)"
        p += [s, s, s]
    if a.get("classe_id"):
        q += " AND s.classe_id=?"; p.append(a.get("classe_id"))
    if a.get("groupe_id"):
        q += " AND s.groupe_id=?"; p.append(a.get("groupe_id"))
    if a.get("sexe"):
        q += " AND s.sexe=?"; p.append(a.get("sexe"))
    q += " ORDER BY s.nom COLLATE NOCASE, s.prenom COLLATE NOCASE"
    c = connect()
    rows = c.execute(q, p).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/students", methods=["POST"])
def add_student():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    nom = (d.get("nom") or "").strip()
    prenom = (d.get("prenom") or "").strip()
    if not nom or not prenom:
        return jsonify({"error": "Mila anarana sy fanampiny"}), 400
    matricule = (d.get("matricule") or "").strip() or gen_matricule()
    age = d.get("age") or compute_age(d.get("date_naissance"))
    c = connect()
    if c.execute("SELECT id FROM students WHERE matricule=?", (matricule,)).fetchone():
        c.close(); return jsonify({"error": "Efa misy io matricule io"}), 400
    cur = c.execute("""INSERT INTO students
        (matricule, nom, prenom, date_naissance, age, sexe, lieu, telephone,
         parent_nom, parent_telephone, classe_id, groupe_id, situation, photo)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (matricule, nom, prenom, d.get("date_naissance"), age,
         d.get("sexe", "M"), d.get("lieu", ""), d.get("telephone", ""),
         d.get("parent_nom", ""), d.get("parent_telephone", ""),
         d.get("classe_id"), d.get("groupe_id"), d.get("situation", ""), d.get("photo", "")))
    c.commit()
    result = student_full(c, cur.lastrowid)
    c.close()
    return jsonify(result), 201

@app.route("/students/<int:sid>")
def get_student(sid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    s = student_full(c, sid)
    if not s:
        c.close(); return jsonify({"error": "Tsy hita"}), 404
    hist = c.execute("SELECT * FROM attendance WHERE student_id=? ORDER BY date DESC LIMIT 60", (sid,)).fetchall()
    stats = {
        "present": c.execute("SELECT COUNT(*) FROM attendance WHERE student_id=? AND statut='present'", (sid,)).fetchone()[0],
        "absent": c.execute("SELECT COUNT(*) FROM attendance WHERE student_id=? AND statut='absent'", (sid,)).fetchone()[0],
        "retard": c.execute("SELECT COUNT(*) FROM attendance WHERE student_id=? AND statut='retard'", (sid,)).fetchone()[0],
        "excuse": c.execute("SELECT COUNT(*) FROM attendance WHERE student_id=? AND statut='excuse'", (sid,)).fetchone()[0],
    }
    total = sum(stats.values())
    stats["total"] = total
    stats["taux"] = round((stats["present"] + stats["retard"]) / total * 100, 1) if total else 0
    c.close()
    return jsonify({**s, "history": [dict(h) for h in hist], "stats": stats})

@app.route("/students/<int:sid>", methods=["PUT"])
def update_student(sid):
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    fields, vals = [], []
    for k in ["matricule","nom","prenom","date_naissance","age","sexe","lieu",
              "telephone","parent_nom","parent_telephone","classe_id","groupe_id",
              "situation","photo"]:
        if k in d:
            fields.append(f"{k}=?"); vals.append(d[k])
    if "date_naissance" in d and "age" not in d:
        fields.append("age=?"); vals.append(compute_age(d["date_naissance"]))
    fields.append("updated_at=CURRENT_TIMESTAMP")
    vals.append(sid)
    c = connect()
    c.execute(f"UPDATE students SET {', '.join(fields)} WHERE id=?", vals)
    c.commit()
    r = student_full(c, sid)
    c.close()
    return jsonify(r) if r else (jsonify({"error": "Tsy hita"}), 404)

@app.route("/students/<int:sid>", methods=["DELETE"])
def del_student(sid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    c.execute("DELETE FROM attendance WHERE student_id=?", (sid,))
    c.execute("DELETE FROM notes WHERE student_id=?", (sid,))
    c.execute("DELETE FROM students WHERE id=?", (sid,))
    c.commit()
    c.close()
    return jsonify({"ok": True})

@app.route("/students/<int:sid>/photo", methods=["PUT"])
def update_photo(sid):
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    c = connect()
    c.execute("UPDATE students SET photo=? WHERE id=?", (d.get("photo", ""), sid))
    c.commit()
    c.close()
    return jsonify({"ok": True})

@app.route("/students/<int:sid>/qr")
def student_qr(sid):
    c = connect()
    s = c.execute("SELECT matricule FROM students WHERE id=?", (sid,)).fetchone()
    c.close()
    if not s:
        return jsonify({"error": "Tsy hita"}), 404
    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(f"FANIRINTSOA:{s['matricule']}")
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return send_file(buf, mimetype="image/png")

@app.route("/students/<int:sid>/badge")
def student_badge(sid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    s = student_full(c, sid)
    c.close()
    if not s:
        return jsonify({"error": "Tsy hita"}), 404
    return jsonify({
        "id": s["id"], "matricule": s["matricule"],
        "nom": s["nom"], "prenom": s["prenom"],
        "classe": s["classe_nom"], "groupe": s["groupe_nom"],
        "photo": s["photo"],
    })

# ============================================================
# EXCEL IMPORT / EXPORT
# ============================================================
@app.route("/students/template")
def excel_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Eleves"
    ws.append(["matricule","nom","prenom","date_naissance","age","sexe",
               "lieu","parent_nom","parent_telephone","telephone","classe"])
    ws.append(["","Rakoto","Jean","2010-05-15",15,"M","Antananarivo",
               "Rabe","0341234567","0341111111","Classe A"])
    for col, w in zip("ABCDEFGHIJK", [15,15,15,15,6,6,18,18,18,15,12]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name="modele_eleves.xlsx", as_attachment=True)

@app.route("/students/import-excel", methods=["POST"])
def import_excel():
    uid, err = require_auth(request)
    if err: return err
    if "file" not in request.files:
        return jsonify({"error": "Tsy misy fichier"}), 400
    f = request.files["file"]
    if not f.filename.endswith((".xlsx", ".xls")):
        return jsonify({"error": "Mila .xlsx na .xls"}), 400
    try:
        wb = load_workbook(f, data_only=True)
        ws = wb.active
    except Exception as e:
        return jsonify({"error": f"Tsy vakiana: {e}"}), 400
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return jsonify({"error": "Fichier foana"}), 400
    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    col = {h: i for i, h in enumerate(headers)}
    if "nom" not in col or "prenom" not in col:
        return jsonify({"error": "Mila colonnes: nom, prenom"}), 400
    c = connect()
    classes_map = {r["nom"].lower(): r["id"] for r in c.execute("SELECT id, nom FROM classes").fetchall()}
    results = {"total": len(rows) - 1, "imported": 0, "duplicates": 0, "errors": [], "preview": []}
    for idx, row in enumerate(rows[1:], start=2):
        try:
            def v(k):
                i = col.get(k)
                if i is None or i >= len(row): return None
                return row[i]
            nom = str(v("nom") or "").strip()
            prenom = str(v("prenom") or "").strip()
            if not nom or not prenom:
                results["errors"].append({"ligne": idx, "error": "Mila nom sy prenom"})
                continue
            matricule = str(v("matricule") or "").strip() or gen_matricule()
            if c.execute("SELECT id FROM students WHERE matricule=?", (matricule,)).fetchone():
                results["duplicates"] += 1
                continue
            classe_nom = str(v("classe") or "").strip()
            classe_id = classes_map.get(classe_nom.lower()) if classe_nom else None
            if classe_nom and not classe_id:
                cur = c.execute("INSERT INTO classes (nom) VALUES (?)", (classe_nom,))
                classe_id = cur.lastrowid
                classes_map[classe_nom.lower()] = classe_id
            date_n = str(v("date_naissance") or "").strip() or None
            age = v("age") or compute_age(date_n)
            sexe = str(v("sexe") or "M").strip().upper()[:1] or "M"
            c.execute("""INSERT INTO students
                (matricule,nom,prenom,date_naissance,age,sexe,lieu,
                 parent_nom,parent_telephone,telephone,classe_id)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (matricule,nom,prenom,date_n,age,sexe,
                 str(v("lieu") or ""), str(v("parent_nom") or ""),
                 str(v("parent_telephone") or ""), str(v("telephone") or ""), classe_id))
            results["imported"] += 1
            if len(results["preview"]) < 10:
                results["preview"].append({"matricule": matricule, "nom": nom,
                                           "prenom": prenom, "classe": classe_nom})
        except Exception as e:
            results["errors"].append({"ligne": idx, "error": str(e)})
    c.commit()
    c.close()
    return jsonify(results), 201

# ============================================================
# SESSIONS
# ============================================================
@app.route("/sessions")
def list_sessions():
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    rows = c.execute("""SELECT s.*, cl.nom AS classe_nom, g.nom AS groupe_nom
                        FROM attendance_sessions s
                        LEFT JOIN classes cl ON cl.id = s.classe_id
                        LEFT JOIN groupes g ON g.id = s.groupe_id
                        ORDER BY s.date DESC, s.id DESC LIMIT 100""").fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/sessions", methods=["POST"])
def create_session():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    date_ = d.get("date") or datetime.now().strftime("%Y-%m-%d")
    c = connect()
    cur = c.execute("""INSERT INTO attendance_sessions
        (classe_id, groupe_id, date, heure_debut, heure_fin, statut)
        VALUES (?,?,?,?,?,'open')""",
        (d.get("classe_id"), d.get("groupe_id"), date_,
         d.get("heure_debut"), d.get("heure_fin")))
    c.commit()
    row = c.execute("SELECT * FROM attendance_sessions WHERE id=?", (cur.lastrowid,)).fetchone()
    c.close()
    return jsonify(dict(row)), 201

@app.route("/sessions/<int:sid>/close", methods=["POST"])
def close_session(sid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    s = c.execute("SELECT * FROM attendance_sessions WHERE id=?", (sid,)).fetchone()
    if not s:
        c.close(); return jsonify({"error": "Tsy hita"}), 404
    if s["statut"] == "closed":
        c.close(); return jsonify({"error": "Efa mihidy"}), 400
    q = "SELECT id FROM students WHERE 1=1"
    p = []
    if s["classe_id"]: q += " AND classe_id=?"; p.append(s["classe_id"])
    if s["groupe_id"]: q += " AND groupe_id=?"; p.append(s["groupe_id"])
    expected = [r["id"] for r in c.execute(q, p).fetchall()]
    present = {r["student_id"] for r in c.execute(
        "SELECT student_id FROM attendance WHERE session_id=?", (sid,)).fetchall()}
    absent_count = 0
    for stid in expected:
        if stid not in present:
            c.execute("""INSERT INTO attendance
                (student_id, session_id, date, statut, methode)
                VALUES (?,?,?,'absent','auto')""", (stid, sid, s["date"]))
            absent_count += 1
    c.execute("UPDATE attendance_sessions SET statut='closed' WHERE id=?", (sid,))
    c.commit()
    c.close()
    return jsonify({"ok": True, "absents_marques": absent_count})

# ============================================================
# ATTENDANCE
# ============================================================
@app.route("/attendance")
def list_attendance():
    uid, err = require_auth(request)
    if err: return err
    a = request.args
    q = """SELECT a.*, s.matricule, s.nom, s.prenom, s.photo,
                  cl.nom AS classe_nom, g.nom AS groupe_nom
           FROM attendance a
           JOIN students s ON s.id = a.student_id
           LEFT JOIN classes cl ON cl.id = s.classe_id
           LEFT JOIN groupes g ON g.id = s.groupe_id WHERE 1=1"""
    p = []
    if a.get("date"): q += " AND a.date=?"; p.append(a.get("date"))
    if a.get("statut"): q += " AND a.statut=?"; p.append(a.get("statut"))
    if a.get("classe_id"): q += " AND s.classe_id=?"; p.append(a.get("classe_id"))
    if a.get("groupe_id"): q += " AND s.groupe_id=?"; p.append(a.get("groupe_id"))
    if a.get("student_id"): q += " AND a.student_id=?"; p.append(a.get("student_id"))
    q += " ORDER BY a.date DESC, a.id DESC LIMIT 500"
    c = connect()
    rows = c.execute(q, p).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/attendance/today")
def attendance_today():
    uid, err = require_auth(request)
    if err: return err
    today = datetime.now().strftime("%Y-%m-%d")
    c = connect()
    rows = c.execute("""SELECT a.*, s.matricule, s.nom, s.prenom, s.photo,
                               cl.nom AS classe_nom, g.nom AS groupe_nom
                        FROM attendance a
                        JOIN students s ON s.id = a.student_id
                        LEFT JOIN classes cl ON cl.id = s.classe_id
                        LEFT JOIN groupes g ON g.id = s.groupe_id
                        WHERE a.date=? ORDER BY a.id DESC""", (today,)).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/attendance/check-in", methods=["POST"])
def check_in():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    matricule = (d.get("matricule") or "").strip()
    if not matricule:
        return jsonify({"error": "Mila matricule"}), 400
    if matricule.startswith("FANIRINTSOA:"):
        matricule = matricule.split(":", 1)[1]
    c = connect()
    s = c.execute("SELECT * FROM students WHERE matricule=?", (matricule,)).fetchone()
    if not s:
        c.close(); return jsonify({"error": "Tsy hita io mpianatra io"}), 404
    today = datetime.now().strftime("%Y-%m-%d")
    now = datetime.now().strftime("%H:%M:%S")
    existing = c.execute("SELECT id FROM attendance WHERE student_id=? AND date=?",
                         (s["id"], today)).fetchone()
    if existing:
        c.close()
        return jsonify({"error": "Efa voamarika ho présent io mpianatra io"}), 400
    statut = "present"
    if d.get("heure_limite") and now > d["heure_limite"]:
        statut = "retard"
    cur = c.execute("""INSERT INTO attendance
        (student_id, session_id, date, heure_arrivee, statut, methode)
        VALUES (?,?,?,?,?,?)""",
        (s["id"], d.get("session_id"), today, now, statut, d.get("methode", "qr")))
    c.commit()
    c.close()
    return jsonify({"ok": True,
                    "student": {"nom": s["nom"], "prenom": s["prenom"],
                                "matricule": s["matricule"], "photo": s["photo"]},
                    "heure_arrivee": now, "statut": statut}), 201

@app.route("/attendance/scan", methods=["POST"])
def scan_qr():
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    code = (d.get("code") or "").strip()
    heure_limite = d.get("heure_limite", "08:00")
    if not code:
        return jsonify({"error": "Mila code"}), 400
    matricule = code.replace("FANIRINTSOA:", "").strip()
    c = connect()
    s = c.execute("SELECT * FROM students WHERE matricule=?", (matricule,)).fetchone()
    if not s:
        c.close()
        return jsonify({"error": "Tsy hita io mpianatra io", "code": code}), 404
    today = datetime.now().strftime("%Y-%m-%d")
    now = datetime.now().strftime("%H:%M:%S")
    existing = c.execute("SELECT * FROM attendance WHERE student_id=? AND date=?",
                         (s["id"], today)).fetchone()
    if existing:
        c.close()
        return jsonify({
            "error": "Efa voamarika",
            "already": True,
            "student": {"id": s["id"], "nom": s["nom"], "prenom": s["prenom"],
                        "matricule": s["matricule"], "photo": s["photo"]},
            "heure_arrivee": existing["heure_arrivee"],
            "statut": existing["statut"],
        }), 400
    statut = "present"
    if now > heure_limite:
        statut = "retard"
    cur = c.execute("""INSERT INTO attendance
        (student_id, session_id, date, heure_arrivee, statut, methode)
        VALUES (?,?,?,?,?,'qr')""",
        (s["id"], d.get("session_id"), today, now, statut))
    c.commit()
    classe = c.execute("SELECT nom FROM classes WHERE id=?", (s["classe_id"],)).fetchone()
    groupe = c.execute("SELECT nom FROM groupes WHERE id=?", (s["groupe_id"],)).fetchone()
    c.close()
    return jsonify({
        "ok": True,
        "student": {"id": s["id"], "matricule": s["matricule"],
                    "nom": s["nom"], "prenom": s["prenom"], "photo": s["photo"],
                    "classe": classe["nom"] if classe else None,
                    "groupe": groupe["nom"] if groupe else None},
        "attendance": {"id": cur.lastrowid, "date": today,
                       "heure_arrivee": now, "statut": statut},
    }), 201

@app.route("/attendance/<int:aid>", methods=["PUT"])
def update_attendance(aid):
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    fields, vals = [], []
    for k in ["statut", "heure_arrivee", "heure_depart"]:
        if k in d:
            fields.append(f"{k}=?"); vals.append(d[k])
    if not fields: return jsonify({"error": "Tsy misy"}), 400
    vals.append(aid)
    c = connect()
    c.execute(f"UPDATE attendance SET {', '.join(fields)} WHERE id=?", vals)
    c.commit()
    r = c.execute("SELECT * FROM attendance WHERE id=?", (aid,)).fetchone()
    c.close()
    return jsonify(dict(r)) if r else (jsonify({"error": "Tsy hita"}), 404)

@app.route("/attendance/<int:aid>", methods=["DELETE"])
def del_attendance(aid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    c.execute("DELETE FROM attendance WHERE id=?", (aid,))
    c.commit()
    c.close()
    return jsonify({"ok": True})

@app.route("/attendance/<int:aid>/justify", methods=["PUT"])
def justify_absence(aid):
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    motif = (d.get("motif") or "").strip()
    if not motif: return jsonify({"error": "Mila motif"}), 400
    c = connect()
    c.execute("""UPDATE attendance SET statut='excuse',
                 justif_motif=?, justif_date=CURRENT_TIMESTAMP WHERE id=?""", (motif, aid))
    c.commit()
    row = c.execute("SELECT * FROM attendance WHERE id=?", (aid,)).fetchone()
    c.close()
    return jsonify(dict(row)) if row else (jsonify({"error": "Tsy hita"}), 404)

# ============================================================
# DASHBOARD
# ============================================================
@app.route("/dashboard")
def dashboard():
    uid, err = require_auth(request)
    if err: return err
    today = datetime.now().strftime("%Y-%m-%d")
    c = connect()
    total_students = c.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    total_classes = c.execute("SELECT COUNT(*) FROM classes").fetchone()[0]
    total_groupes = c.execute("SELECT COUNT(*) FROM groupes").fetchone()[0]
    present_today = c.execute("SELECT COUNT(*) FROM attendance WHERE date=? AND statut='present'", (today,)).fetchone()[0]
    late_today = c.execute("SELECT COUNT(*) FROM attendance WHERE date=? AND statut='retard'", (today,)).fetchone()[0]
    absent_today = c.execute("SELECT COUNT(*) FROM attendance WHERE date=? AND statut='absent'", (today,)).fetchone()[0]
    excuse_today = c.execute("SELECT COUNT(*) FROM attendance WHERE date=? AND statut='excuse'", (today,)).fetchone()[0]
    total_marked = present_today + late_today + absent_today + excuse_today
    taux = round((present_today + late_today) / total_marked * 100, 1) if total_marked else 0
    evolution = []
    for i in range(6, -1, -1):
        d = (datetime.now() - timedelta(days=i)).strftime("%Y-%m-%d")
        p = c.execute("SELECT COUNT(*) FROM attendance WHERE date=? AND statut IN ('present','retard')", (d,)).fetchone()[0]
        ab = c.execute("SELECT COUNT(*) FROM attendance WHERE date=? AND statut='absent'", (d,)).fetchone()[0]
        evolution.append({"date": d, "present": p, "absent": ab})
    c.close()
    return jsonify({
        "total_students": total_students, "total_classes": total_classes,
        "total_groupes": total_groupes,
        "present_today": present_today, "late_today": late_today,
        "absent_today": absent_today, "excuse_today": excuse_today,
        "taux_presence": taux, "evolution": evolution,
    })

# ============================================================
# STATISTIQUES PRÉSENCE
# ============================================================
@app.route("/attendance/statistics")
def stats_attendance():
    uid, err = require_auth(request)
    if err: return err
    a = request.args
    base = "FROM attendance a JOIN students s ON s.id = a.student_id WHERE 1=1"
    params = []
    for k, col in [("from","a.date>="), ("to","a.date<="), ("classe_id","s.classe_id="),
                   ("groupe_id","s.groupe_id="), ("student_id","s.id=")]:
        if a.get(k):
            base += f" AND {col}?"; params.append(a.get(k))
    c = connect()
    totals = c.execute(f"""SELECT COUNT(*) AS total,
        SUM(CASE WHEN a.statut='present' THEN 1 ELSE 0 END) AS present,
        SUM(CASE WHEN a.statut='absent' THEN 1 ELSE 0 END) AS absent,
        SUM(CASE WHEN a.statut='retard' THEN 1 ELSE 0 END) AS retard,
        SUM(CASE WHEN a.statut='excuse' THEN 1 ELSE 0 END) AS excuse
        {base}""", params).fetchone()
    totals = dict(totals)
    t = totals["total"] or 1
    totals["taux"] = round((totals["present"] + totals["retard"]) / t * 100, 1)
    daily = c.execute(f"""SELECT a.date,
        SUM(CASE WHEN a.statut IN ('present','retard') THEN 1 ELSE 0 END) AS present,
        SUM(CASE WHEN a.statut='absent' THEN 1 ELSE 0 END) AS absent
        {base} GROUP BY a.date ORDER BY a.date DESC LIMIT 30""", params).fetchall()
    top = c.execute(f"""SELECT s.id, s.nom, s.prenom, s.matricule,
        SUM(CASE WHEN a.statut IN ('present','retard') THEN 1 ELSE 0 END) AS presences,
        SUM(CASE WHEN a.statut='absent' THEN 1 ELSE 0 END) AS absences
        {base} GROUP BY s.id ORDER BY presences DESC LIMIT 10""", params).fetchall()
    c.close()
    return jsonify({"totals": totals,
                    "daily": [dict(d) for d in daily],
                    "top": [dict(t) for t in top]})

@app.route("/attendance/student/<int:sid>")
def student_history(sid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    rows = c.execute("SELECT * FROM attendance WHERE student_id=? ORDER BY date DESC LIMIT 100",
                     (sid,)).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/attendance/calendar")
def attendance_calendar():
    uid, err = require_auth(request)
    if err: return err
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    c = connect()
    rows = c.execute("""SELECT date, COUNT(*) AS total,
        SUM(CASE WHEN statut IN ('present','retard') THEN 1 ELSE 0 END) AS present,
        SUM(CASE WHEN statut='absent' THEN 1 ELSE 0 END) AS absent
        FROM attendance WHERE date LIKE ?
        GROUP BY date""", (f"{month}%",)).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

# ============================================================
# MATIERES
# ============================================================
@app.route("/matieres")
def list_matieres():
    uid, err = require_auth(request)
    if err: return err
    classe_id = request.args.get("classe_id")
    c = connect()
    if classe_id:
        rows = c.execute("SELECT * FROM matieres WHERE classe_id=? OR classe_id IS NULL ORDER BY nom", (classe_id,)).fetchall()
    else:
        rows = c.execute("SELECT * FROM matieres ORDER BY nom").fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/matieres", methods=["POST"])
def add_matiere():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    n = (d.get("nom") or "").strip()
    if not n: return jsonify({"error": "Mila anarana"}), 400
    c = connect()
    cur = c.execute("INSERT INTO matieres (nom, coefficient, classe_id) VALUES (?,?,?)",
                    (n, d.get("coefficient", 1), d.get("classe_id")))
    c.commit()
    row = c.execute("SELECT * FROM matieres WHERE id=?", (cur.lastrowid,)).fetchone()
    c.close()
    return jsonify(dict(row)), 201

@app.route("/matieres/<int:mid>", methods=["PUT"])
def update_matiere(mid):
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    fields, vals = [], []
    for k in ["nom", "coefficient", "classe_id"]:
        if k in d:
            fields.append(f"{k}=?"); vals.append(d[k])
    if not fields: return jsonify({"error": "Tsy misy"}), 400
    vals.append(mid)
    c = connect()
    c.execute(f"UPDATE matieres SET {', '.join(fields)} WHERE id=?", vals)
    c.commit()
    row = c.execute("SELECT * FROM matieres WHERE id=?", (mid,)).fetchone()
    c.close()
    return jsonify(dict(row)) if row else (jsonify({"error": "Tsy hita"}), 404)

@app.route("/matieres/<int:mid>", methods=["DELETE"])
def del_matiere(mid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    c.execute("DELETE FROM notes WHERE matiere_id=?", (mid,))
    c.execute("DELETE FROM matieres WHERE id=?", (mid,))
    c.commit()
    c.close()
    return jsonify({"ok": True})

# ============================================================
# NOTES
# ============================================================
@app.route("/notes")
def list_notes():
    uid, err = require_auth(request)
    if err: return err
    a = request.args
    q = """SELECT n.*, s.nom, s.prenom, s.matricule, m.nom AS matiere_nom,
                  m.coefficient, cl.nom AS classe_nom
           FROM notes n
           JOIN students s ON s.id = n.student_id
           JOIN matieres m ON m.id = n.matiere_id
           LEFT JOIN classes cl ON cl.id = s.classe_id
           WHERE 1=1"""
    p = []
    if a.get("student_id"): q += " AND n.student_id=?"; p.append(a.get("student_id"))
    if a.get("matiere_id"): q += " AND n.matiere_id=?"; p.append(a.get("matiere_id"))
    if a.get("trimestre"): q += " AND n.trimestre=?"; p.append(a.get("trimestre"))
    if a.get("classe_id"): q += " AND s.classe_id=?"; p.append(a.get("classe_id"))
    q += " ORDER BY n.date DESC, n.id DESC LIMIT 500"
    c = connect()
    rows = c.execute(q, p).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

@app.route("/notes", methods=["POST"])
def add_note():
    uid, err = require_auth(request)
    if err: return err
    d = request.json
    if not d.get("student_id") or not d.get("matiere_id"):
        return jsonify({"error": "Mila student_id sy matiere_id"}), 400
    note = d.get("note")
    if note is None or note == "":
        return jsonify({"error": "Mila note"}), 400
    try:
        note = float(note)
    except:
        return jsonify({"error": "Tsy mety ny note"}), 400
    if note < 0 or note > 20:
        return jsonify({"error": "0 ka hatramin'ny 20 ny note"}), 400
    c = connect()
    cur = c.execute("""INSERT INTO notes
        (student_id, matiere_id, note, type, trimestre, date, remarque)
        VALUES (?,?,?,?,?,?,?)""",
        (d["student_id"], d["matiere_id"], note, d.get("type", "devoir"),
         d.get("trimestre", 1), d.get("date") or datetime.now().strftime("%Y-%m-%d"),
         d.get("remarque", "")))
    c.commit()
    row = c.execute("""SELECT n.*, s.nom, s.prenom, m.nom AS matiere_nom
                       FROM notes n
                       JOIN students s ON s.id = n.student_id
                       JOIN matieres m ON m.id = n.matiere_id
                       WHERE n.id=?""", (cur.lastrowid,)).fetchone()
    c.close()
    return jsonify(dict(row)), 201

@app.route("/notes/<int:nid>", methods=["PUT"])
def update_note(nid):
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    fields, vals = [], []
    for k in ["note", "type", "trimestre", "date", "remarque"]:
        if k in d:
            fields.append(f"{k}=?"); vals.append(d[k])
    if not fields: return jsonify({"error": "Tsy misy"}), 400
    vals.append(nid)
    c = connect()
    c.execute(f"UPDATE notes SET {', '.join(fields)} WHERE id=?", vals)
    c.commit()
    row = c.execute("SELECT * FROM notes WHERE id=?", (nid,)).fetchone()
    c.close()
    return jsonify(dict(row)) if row else (jsonify({"error": "Tsy hita"}), 404)

@app.route("/notes/<int:nid>", methods=["DELETE"])
def del_note(nid):
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    c.execute("DELETE FROM notes WHERE id=?", (nid,))
    c.commit()
    c.close()
    return jsonify({"ok": True})

@app.route("/notes/statistics")
def notes_stats():
    uid, err = require_auth(request)
    if err: return err
    a = request.args
    c = connect()
    q = """SELECT m.nom AS matiere, AVG(n.note) AS moyenne, COUNT(n.id) AS total
           FROM notes n JOIN matieres m ON m.id = n.matiere_id
           WHERE 1=1"""
    p = []
    if a.get("trimestre"): q += " AND n.trimestre=?"; p.append(a.get("trimestre"))
    if a.get("classe_id"): q += " AND n.student_id IN (SELECT id FROM students WHERE classe_id=?)"; p.append(a.get("classe_id"))
    q += " GROUP BY m.id ORDER BY m.nom"
    rows = c.execute(q, p).fetchall()
    c.close()
    return jsonify([dict(r) for r in rows])

# ============================================================
# BULLETIN
# ============================================================
@app.route("/students/<int:sid>/bulletin")
def bulletin(sid):
    uid, err = require_auth(request)
    if err: return err
    trimestre = request.args.get("trimestre", 1)
    c = connect()
    s = c.execute("""SELECT s.*, cl.nom AS classe_nom
                     FROM students s
                     LEFT JOIN classes cl ON cl.id = s.classe_id
                     WHERE s.id=?""", (sid,)).fetchone()
    if not s:
        c.close(); return jsonify({"error": "Tsy hita"}), 404
    rows = c.execute("""SELECT m.id AS matiere_id, m.nom AS matiere_nom,
                               m.coefficient,
                               AVG(n.note) AS moyenne,
                               COUNT(n.id) AS nb_notes
                        FROM matieres m
                        LEFT JOIN notes n ON n.matiere_id = m.id
                            AND n.student_id = ? AND n.trimestre = ?
                        GROUP BY m.id
                        ORDER BY m.nom""", (sid, trimestre)).fetchall()
    matieres = []
    total_pts = 0
    total_coef = 0
    for r in rows:
        d = dict(r)
        if d["moyenne"] is not None:
            total_pts += d["moyenne"] * d["coefficient"]
            total_coef += d["coefficient"]
        matieres.append(d)
    moyenne_generale = round(total_pts / total_coef, 2) if total_coef > 0 else 0
    c.close()
    return jsonify({
        "student": dict(s),
        "trimestre": int(trimestre),
        "matieres": matieres,
        "moyenne_generale": moyenne_generale,
        "total_coefficients": total_coef,
    })

@app.route("/classes/<int:cid>/bulletin")
def class_bulletin(cid):
    uid, err = require_auth(request)
    if err: return err
    trimestre = request.args.get("trimestre", 1)
    c = connect()
    students = c.execute("SELECT * FROM students WHERE classe_id=? ORDER BY nom", (cid,)).fetchall()
    result = []
    for s in students:
        rows = c.execute("""SELECT AVG(n.note) AS moyenne, m.coefficient
                            FROM notes n JOIN matieres m ON m.id = n.matiere_id
                            WHERE n.student_id=? AND n.trimestre=?""", (s["id"], trimestre)).fetchall()
        total_pts = sum(r["moyenne"] * r["coefficient"] for r in rows if r["moyenne"] is not None)
        total_coef = sum(r["coefficient"] for r in rows if r["moyenne"] is not None)
        moy = round(total_pts / total_coef, 2) if total_coef > 0 else 0
        result.append({**dict(s), "moyenne": moy})
    result.sort(key=lambda x: -x["moyenne"])
    for i, r in enumerate(result):
        r["rang"] = i + 1
    c.close()
    return jsonify(result)

# ============================================================
# BACKUP / RESTORE
# ============================================================
@app.route("/backup")
def backup():
    uid, err = require_auth(request)
    if err: return err
    c = connect()
    data = {
        "version": "2.0", "date": datetime.now().isoformat(),
        "classes": [dict(r) for r in c.execute("SELECT * FROM classes").fetchall()],
        "groupes": [dict(r) for r in c.execute("SELECT * FROM groupes").fetchall()],
        "students": [dict(r) for r in c.execute("SELECT * FROM students").fetchall()],
        "attendance": [dict(r) for r in c.execute("SELECT * FROM attendance").fetchall()],
        "sessions": [dict(r) for r in c.execute("SELECT * FROM attendance_sessions").fetchall()],
        "matieres": [dict(r) for r in c.execute("SELECT * FROM matieres").fetchall()],
        "notes": [dict(r) for r in c.execute("SELECT * FROM notes").fetchall()],
    }
    c.close()
    return jsonify(data)

@app.route("/restore", methods=["POST"])
def restore():
    uid, err = require_auth(request)
    if err: return err
    d = request.json or {}
    if not d.get("students"):
        return jsonify({"error": "Fichier foana"}), 400
    c = connect()
    count = 0
    for s in d.get("students", []):
        try:
            c.execute("""INSERT OR IGNORE INTO students
                (id, matricule, nom, prenom, date_naissance, age, sexe, lieu,
                 telephone, parent_nom, parent_telephone, classe_id, groupe_id,
                 situation, photo)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (s.get("id"), s.get("matricule"), s.get("nom"), s.get("prenom"),
                 s.get("date_naissance"), s.get("age"), s.get("sexe"),
                 s.get("lieu"), s.get("telephone"), s.get("parent_nom"),
                 s.get("parent_telephone"), s.get("classe_id"), s.get("groupe_id"),
                 s.get("situation"), s.get("photo", "")))
            count += 1
        except Exception:
            pass
    for m in d.get("matieres", []):
        try:
            c.execute("INSERT OR IGNORE INTO matieres (nom, coefficient, classe_id) VALUES (?,?,?)",
                      (m.get("nom"), m.get("coefficient", 1), m.get("classe_id")))
        except Exception:
            pass
    c.commit()
    c.close()
    return jsonify({"ok": True, "restored": count})

# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port)

